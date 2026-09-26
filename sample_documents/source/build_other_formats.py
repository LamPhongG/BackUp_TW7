"""Build the DOCX, TXT and MD sample documents, and copy the CSV ones, into sample_documents/.

A Markdown source with `format: docx | txt | md` in its front matter becomes
../DOC-XX_<family>_v<version>.<format>; a DOC-XX_<family>.csv source is copied as
../DOC-XX_<family>_v1.0.csv. PDF sources are built by build_pdfs.mjs.

    cd sample_documents/source
    ../../backend/.venv/Scripts/python build_other_formats.py            # all
    ../../backend/.venv/Scripts/python build_other_formats.py DOC-21     # one document
"""
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
COMPANY = "FourAngryBirds EdTech & HR Solutions JSC"
META_FIELDS = [
    ("document_id", "Document ID"), ("title_vi", "Vietnamese title"), ("category", "Category"),
    ("department", "Department"), ("version", "Version"), ("status", "Status"),
    ("effective_date", "Effective date"), ("owner", "Owner"),
]


def parse(path: Path) -> tuple[dict, str]:
    text = path.read_text(encoding="utf-8")
    match = re.match(r"^---\r?\n(.*?)\r?\n---\r?\n", text, re.S)
    meta = dict(line.split(":", 1) for line in match.group(1).splitlines() if ":" in line)
    return {k.strip(): v.strip() for k, v in meta.items()}, text[match.end():]


def add_runs(paragraph, text: str) -> None:
    # **bold** is the only inline markup the sources use.
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = i % 2 == 1


def build_docx(meta: dict, body: str, out: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(22)
    section.left_margin = section.right_margin = Mm(20)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10.5)
    for level, size in ((1, 20), (2, 13), (3, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(0x31, 0x2E, 0x81)
    if "Company Line" not in doc.styles:
        company = doc.styles.add_style("Company Line", WD_STYLE_TYPE.PARAGRAPH)
        company.base_style = normal
        company.font.size, company.font.color.rgb = Pt(9), RGBColor(0x6D, 0x28, 0xD9)

    doc.add_paragraph(COMPANY, style="Company Line")
    for line in body.splitlines():
        line = line.rstrip()
        if not line:
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
            if len(heading.group(1)) == 1:
                rows = [(label, meta[key]) for key, label in META_FIELDS if meta.get(key)]
                table = doc.add_table(rows=len(rows), cols=2, style="Table Grid")
                for row, (label, value) in zip(table.rows, rows, strict=True):
                    row.cells[0].text, row.cells[1].text = label, value
                doc.add_paragraph()
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_runs(doc.add_paragraph(), line)
    doc.save(out)


def build_txt(meta: dict, body: str, out: Path) -> None:
    lines = []
    for line in body.splitlines():
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", heading.group(2) if heading else line)
        lines.append(text)
        if heading and len(heading.group(1)) == 1:
            lines.append(COMPANY)
            lines += [f"{label}: {meta[key]}" for key, label in META_FIELDS if meta.get(key)]
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def build_md(path: Path, out: Path) -> None:
    text = path.read_text(encoding="utf-8")
    out.write_text(re.sub(r"(?m)^format: .*\n", "", text, count=1), encoding="utf-8")


def main(only: set[str]) -> None:
    for path in sorted(HERE.glob("DOC-*.md")):
        meta, body = parse(path)
        fmt = meta.get("format", "pdf")
        if fmt == "pdf" or (only and path.name[:6] not in only):
            continue
        out = OUT / f"{re.sub(r'(_v[\d.]+)?$', '', path.stem)}_v{meta.get('version', '1.0')}.{fmt}"
        if fmt == "docx":
            build_docx(meta, body, out)
        elif fmt == "txt":
            build_txt(meta, body, out)
        else:
            build_md(path, out)
        print("wrote", out.name)
    for path in sorted(HERE.glob("DOC-*.csv")):
        if only and path.name[:6] not in only:
            continue
        out = OUT / f"{path.stem}_v1.0.csv"
        shutil.copyfile(path, out)
        print("wrote", out.name)


if __name__ == "__main__":
    main(set(sys.argv[1:]))
