# Unit tests for document processing and validation pipeline

from pathlib import Path

import pytest

from src.document_processing.chunker import split_into_chunks
from src.document_validation.validator import (
    FileSizeError,
    UnsupportedFormatError,
    validate_document,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_raw_page(
    text: str,
    page_number: int = 1,
    doc_id: str = "abc123def456",
    source_file: str = "policy.pdf",
) -> dict:
    """Builds a minimal raw-page dict matching the reader output contract."""
    return {
        "doc_id": doc_id,
        "page_number": page_number,
        "raw_text": text,
        "source_file": source_file,
    }


# ---------------------------------------------------------------------------
# Validator tests
# ---------------------------------------------------------------------------

class TestValidateDocument:
    def test_raises_when_file_missing(self, tmp_path):
        missing = tmp_path / "ghost.pdf"
        with pytest.raises(FileNotFoundError):
            validate_document(missing)

    def test_raises_for_unsupported_extension(self, tmp_path):
        txt_file = tmp_path / "notes.txt"
        txt_file.write_bytes(b"x" * 1024)
        with pytest.raises(UnsupportedFormatError):
            validate_document(txt_file)

    def test_raises_for_oversized_file(self, tmp_path):
        big_file = tmp_path / "huge.pdf"
        # Write slightly over 50 MB
        big_file.write_bytes(b"x" * (51 * 1024 * 1024))
        with pytest.raises(FileSizeError, match="50MB"):
            validate_document(big_file)

    def test_raises_for_empty_file(self, tmp_path):
        empty = tmp_path / "empty.docx"
        empty.write_bytes(b"")
        with pytest.raises(FileSizeError, match="bytes"):
            validate_document(empty)

    def test_passes_for_valid_pdf(self, tmp_path):
        valid_pdf = tmp_path / "policy.pdf"
        valid_pdf.write_bytes(b"x" * 1024)
        # Should not raise — return value is None (gate, not transformer)
        assert validate_document(valid_pdf) is None

    def test_passes_for_valid_docx(self, tmp_path):
        valid_docx = tmp_path / "handbook.docx"
        valid_docx.write_bytes(b"x" * 2048)
        assert validate_document(valid_docx) is None


# ---------------------------------------------------------------------------
# Chunker tests
# ---------------------------------------------------------------------------

class TestSplitIntoChunks:
    def test_raises_on_empty_input(self):
        with pytest.raises(ValueError, match="empty"):
            split_into_chunks([])

    def test_raises_on_missing_required_key(self):
        bad_page = {"doc_id": "abc", "page_number": 1, "raw_text": "text"}
        # source_file is missing
        with pytest.raises(ValueError, match="source_file"):
            split_into_chunks([bad_page])

    def test_single_page_produces_at_least_one_chunk(self):
        # Simulate one page of dense policy text with no headings
        text = "All employees must complete compliance training within 30 days. " * 8
        pages = [_make_raw_page(text)]
        chunks = split_into_chunks(pages)
        assert len(chunks) >= 1

    def test_all_caps_line_triggers_new_section(self):
        preamble = "Some preamble text that is long enough to form a valid chunk. " * 3
        heading_section = "Employees are entitled to 12 days of annual leave per calendar year. " * 3
        text = preamble + "\nLEAVE ENTITLEMENT\n" + heading_section
        pages = [_make_raw_page(text)]
        chunks = split_into_chunks(pages)
        headings = [c.heading for c in chunks]
        assert "LEAVE ENTITLEMENT" in headings


    def test_chunk_inherits_correct_page_number(self):
        page1 = _make_raw_page(
            "HEALTH AND SAFETY POLICY\n" + "Safety is paramount. " * 10,
            page_number=3
        )
        chunks = split_into_chunks([page1])
        # The chunk for the heading on page 3 must carry page 3
        safety_chunks = [c for c in chunks if "HEALTH" in c.heading]
        assert all(c.page_number == 3 for c in safety_chunks)

    def test_chunk_ids_are_unique(self):
        text = "SECTION ONE\n" + "Content A. " * 10 + "\nSECTION TWO\n" + "Content B. " * 10
        pages = [_make_raw_page(text)]
        chunks = split_into_chunks(pages)
        ids = [c.chunk_id for c in chunks]
        assert len(ids) == len(set(ids)), "Duplicate chunk_ids detected"

    def test_chunk_content_does_not_exceed_max_length(self):
        # Oversized content must be truncated, not dropped
        huge_text = "Policy detail. " * 300  # well over 2000 chars
        pages = [_make_raw_page(huge_text)]
        chunks = split_into_chunks(pages)
        for chunk in chunks:
            assert len(chunk.content) <= 2000

    def test_doc_id_is_propagated_to_all_chunks(self):
        doc_id = "testdoc000abc"
        text = "SECTION ONE\n" + "Content. " * 20
        pages = [_make_raw_page(text, doc_id=doc_id)]
        chunks = split_into_chunks(pages)
        assert all(c.doc_id == doc_id for c in chunks)


# ---------------------------------------------------------------------------
# Reader hardening & edge cases tests
# ---------------------------------------------------------------------------

class TestHardenedDocumentReaders:
    def test_raises_for_scanned_image_only_pdf(self, tmp_path):
        import fitz
        from src.document_processing.pdf_reader import PDFReadError, read_pdf

        pdf_path = tmp_path / "scanned.pdf"
        doc = fitz.open()
        doc.new_page()  # Blank page without any selectable text
        doc.save(str(pdf_path))
        doc.close()

        with pytest.raises(PDFReadError, match="scanned image-only"):
            read_pdf(pdf_path)

    def test_raises_for_encrypted_pdf(self, tmp_path):
        import fitz
        from src.document_processing.pdf_reader import PDFReadError, read_pdf

        pdf_path = tmp_path / "encrypted.pdf"
        doc = fitz.open()
        p = doc.new_page()
        p.insert_text((50, 50), "Confidential data")
        doc.save(str(pdf_path), encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret123")
        doc.close()

        with pytest.raises(PDFReadError, match="encrypted"):
            read_pdf(pdf_path)

    def test_docx_extracts_table_rows(self, tmp_path):
        from docx import Document
        from src.document_processing.docx_reader import read_docx

        docx_path = tmp_path / "policy_table.docx"
        doc = Document()
        doc.add_paragraph("POLICY TABLE SUMMARY")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Role"
        table.cell(0, 1).text = "Annual Leave"
        table.cell(1, 0).text = "Software Engineer"
        table.cell(1, 1).text = "15 Days"
        doc.save(str(docx_path))

        pages = read_docx(docx_path)
        assert len(pages) == 1
        assert "Software Engineer | 15 Days" in pages[0]["raw_text"]

    def test_docx_raises_on_whitespace_only(self, tmp_path):
        from docx import Document
        from src.document_processing.docx_reader import DOCXReadError, read_docx

        docx_path = tmp_path / "blank.docx"
        doc = Document()
        doc.add_paragraph("   \n\t  ")
        doc.save(str(docx_path))

        with pytest.raises(DOCXReadError, match="no readable text content"):
            read_docx(docx_path)

